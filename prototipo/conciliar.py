"""Prototipo de conciliación de pagos del grupo scout.

Cruza el Excel de movimientos que descarga Unicaja con el listado de
inscritos (padres / niño) y sugiere a qué familia pertenece cada ingreso,
con un nivel de confianza.

Uso:
    python conciliar.py MOVIMIENTOS.xls INSCRITOS.(xlsx|xls|csv) [-o salida.xlsx]

Todo funciona en local: no se envía nada a ningún servidor.
"""
import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.styles import Alignment, Font, PatternFill

# Palabras que unen apellidos compuestos: "Pérez de Arce", "Ruiz de la Torre".
ENLACES = {"DE", "DEL", "LA", "LAS", "LOS"}

# Segundos nombres habituales: en "Juan José Rodríguez" JOSE es nombre, no apellido.
SEGUNDOS_NOMBRES = {
    "MARIA", "JOSE", "CARMEN", "ANGEL", "LUIS", "JESUS", "ISABEL", "TERESA",
    "DOLORES", "JOSEFA", "LUISA", "ANGELES", "MANUEL", "ANTONIO", "IGNACIO",
    "JAVIER", "CARLOS", "FRANCISCO", "ELENA", "CRISTINA", "VICTORIA", "EUGENIA",
}

# En el listado se escribe "Pepa" y en el banco sale "JOSEFA".
DIMINUTIVOS = {
    "MARI": "MARIA", "PEPA": "JOSEFA", "PEPE": "JOSE", "PACO": "FRANCISCO",
    "PACA": "FRANCISCA", "LOLA": "DOLORES", "NACHO": "IGNACIO", "MANOLO": "MANUEL",
    "MANOLI": "MANUELA", "CHARO": "ROSARIO", "CONCHA": "CONCEPCION", "MERCHE": "MERCEDES",
    "CHUS": "JESUS", "NATI": "NATIVIDAD", "INMA": "INMACULADA", "RAFA": "RAFAEL",
}

CABECERA_PADRES = r"PADRE|MADRE|TUTOR|FAMILIA"
CABECERA_NINO = r"NINO|NINA|EDUCANDO|SCOUT|HIJO|CHAVAL|NOMBRE"

NIVELES = ["CONFIRMADA", "ALTA", "MEDIA", "BAJA", "AMBIGUA", "SIN IDENTIFICAR"]
COLORES = {
    "CONFIRMADA": "B7E1CD", "ALTA": "C6EFCE", "MEDIA": "FFEB9C",
    "BAJA": "FCE4D6", "AMBIGUA": "F8CBAD", "SIN IDENTIFICAR": "D9D9D9",
}

FICHERO_CONOCIDOS = "pagadores_conocidos.csv"


# ---------------------------------------------------------------- nombres

def normalizar(texto):
    """'Mº Carmen Pérez' -> ['MARIA', 'CARMEN', 'PEREZ'] (sin tildes, en mayúsculas)."""
    texto = re.sub(r"\bM\s*[ºª.]\s*", "MARIA ", str(texto or ""), flags=re.I)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z ]+", " ", texto).upper().split()


def tokens(texto):
    """Como normalizar, pero juntando apellidos compuestos en un solo token."""
    palabras = normalizar(texto)
    resultado, i = [], 0
    while i < len(palabras):
        if palabras[i] in ENLACES:
            j = i
            while j < len(palabras) and palabras[j] in ENLACES:
                j += 1
            if resultado and j < len(palabras):
                resultado[-1] = " ".join([resultado[-1], *palabras[i:j + 1]])
                i = j + 1
            else:
                i = j  # enlace suelto al principio o al final: se ignora
            continue
        resultado.append(DIMINUTIVOS.get(palabras[i], palabras[i]))
        i += 1
    return resultado


def distancia(a, b):
    previa = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        actual = [i]
        for j, cb in enumerate(b, 1):
            actual.append(min(previa[j] + 1, actual[j - 1] + 1, previa[j - 1] + (ca != cb)))
        previa = actual
    return previa[-1]


def parecidos(a, b):
    """Iguales, o con una sola letra de diferencia en palabras largas (erratas)."""
    if a == b:
        return True
    if min(len(a), len(b)) < 5 or abs(len(a) - len(b)) > 1:
        return False
    return distancia(a, b) <= 1


def contiene(token, conjunto):
    return any(parecidos(token, otro) for otro in conjunto)


@dataclass
class Persona:
    original: str
    nombres: list
    apellidos: list
    seccion: int = 0                              # solo niños: bloque del listado
    marcas: list = field(default_factory=list)    # solo niños: actividades ya marcadas

    @property
    def palabras(self):
        return self.nombres + self.apellidos


def persona(texto):
    """Asume el orden del listado de inscritos: Nombre [Nombre2] Apellido1 [Apellido2]."""
    t = tokens(texto)
    if len(t) >= 3:
        n = 2 if t[1] in SEGUNDOS_NOMBRES else 1
        return Persona(texto.strip(), t[:n], t[n:])
    if len(t) == 2:
        return Persona(texto.strip(), t[:1], t[1:])
    return Persona(texto.strip(), t, [])


# ---------------------------------------------------------------- datos

@dataclass
class Familia:
    clave: str
    padres: list = field(default_factory=list)   # [Persona]
    ninos: list = field(default_factory=list)    # [Persona]

    @property
    def etiqueta(self):
        if self.padres:
            return " / ".join(p.original for p in self.padres)
        return "(sin padres) " + ", ".join(n.original for n in self.ninos)

    @property
    def nombres_ninos(self):
        return ", ".join(n.original for n in self.ninos)

    @property
    def secciones(self):
        return ", ".join(str(s) for s in sorted({n.seccion for n in self.ninos}))

    @property
    def marcas(self):
        return "\n".join(f"{n.original}: {', '.join(n.marcas)}" for n in self.ninos if n.marcas)


@dataclass
class Movimiento:
    fecha: object
    concepto: str
    importe: float
    numero: object


def leer_filas(ruta):
    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo == ".csv":
        texto = ruta.read_text(encoding="utf-8-sig")
        separador = ";" if texto.count(";") >= texto.count(",") else ","
        return [fila for fila in csv.reader(texto.splitlines(), delimiter=separador)]
    if sufijo == ".xls":
        hoja = xlrd.open_workbook(ruta).sheet_by_index(0)
        return [hoja.row_values(f) for f in range(hoja.nrows)]
    hoja = openpyxl.load_workbook(ruta, data_only=True, read_only=True).active
    return [["" if c is None else c for c in fila] for fila in hoja.iter_rows(values_only=True)]


def columna(letra_o_numero):
    if letra_o_numero is None:
        return None
    if str(letra_o_numero).isdigit():
        return int(letra_o_numero) - 1
    return ord(str(letra_o_numero).upper()) - ord("A")


def compatibles(x, y):
    """'Mercedes Sánchez' y 'Mercedes Sánchez González' pueden ser la misma persona."""
    corta, larga = sorted((x.palabras, y.palabras), key=len)
    return bool(corta) and parecidos(corta[0], larga[0]) and all(contiene(t, larga) for t in corta)


def misma_familia(familia, padres, nino):
    """Mismos padres (admitiendo erratas o apellidos de menos) y hermanos que comparten apellido."""
    if not familia.padres or not padres:
        return False
    pocos, muchos = sorted((familia.padres, padres), key=len)
    if not all(any(compatibles(p, q) for q in muchos) for p in pocos):
        return False
    apellidos = {a for n in familia.ninos for a in n.apellidos}
    return any(contiene(a, apellidos) for a in nino.apellidos)


def leer_inscritos(ruta, col_padres=None, col_nino=None):
    """Lee el listado. Cada cabecera repetida ('Nombres padres | Educando | Pago...')
    empieza una sección nueva; las columnas a la derecha del niño son actividades."""
    filas = leer_filas(ruta)
    familias, actividades = [], {}
    seccion, ninos_en_seccion = 1, 0
    for fila in filas:
        celdas = [" ".join(normalizar(c)) for c in fila]
        i_padres = [i for i, c in enumerate(celdas) if re.search(CABECERA_PADRES, c)]
        i_ninos = [i for i, c in enumerate(celdas) if re.search(CABECERA_NINO, c) and i not in i_padres]
        if i_padres and i_ninos:
            col_padres = i_padres[0] if col_padres is None else col_padres
            col_nino = i_ninos[0] if col_nino is None else col_nino
            actividades = {i: re.sub(r"^\s*pago( del?)?\s+", "", str(v).strip(), flags=re.I).capitalize()
                           for i, v in enumerate(fila)
                           if i > max(col_padres, col_nino) and str(v or "").strip()}
            if ninos_en_seccion:
                seccion, ninos_en_seccion = seccion + 1, 0
            continue

        cp = 0 if col_padres is None else col_padres
        cn = 1 if col_nino is None else col_nino
        if max(cp, cn) >= len(fila):
            continue
        nino = persona(str(fila[cn] or ""))
        if not nino.palabras:          # filas vacías, "?", espacios...
            continue
        nino.seccion = seccion
        nino.marcas = [nombre for i, nombre in actividades.items()
                       if i < len(fila) and str(fila[i] or "").strip()]
        ninos_en_seccion += 1

        padres = [persona(p) for p in str(fila[cp] or "").split("/") if tokens(p)]
        familia = next((f for f in familias if misma_familia(f, padres, nino)), None)
        if familia is None:
            familia = Familia(f"F{len(familias) + 1:03d}", padres)
            familias.append(familia)
        elif sum(len(p.palabras) for p in padres) > sum(len(p.palabras) for p in familia.padres):
            familia.padres = padres    # nos quedamos con la versión más completa de los nombres
        familia.ninos.append(nino)
    return familias


def leer_movimientos(ruta):
    libro = xlrd.open_workbook(ruta)
    hoja = libro.sheet_by_index(0)
    for f in range(hoja.nrows):
        cabecera = [" ".join(normalizar(c)) for c in hoja.row_values(f)]
        if "CONCEPTO" in cabecera and "IMPORTE" in cabecera:
            break
    else:
        sys.exit("No encuentro la fila de cabecera (Concepto / Importe) en el Excel del banco.")
    i_concepto, i_importe = cabecera.index("CONCEPTO"), cabecera.index("IMPORTE")
    i_fecha = next((i for i, c in enumerate(cabecera) if c.startswith("FECHA")), None)
    i_numero = next((i for i, c in enumerate(cabecera) if c.endswith("MOV")), None)

    movimientos = []
    for f in range(f + 1, hoja.nrows):
        fila = hoja.row_values(f)
        if not isinstance(fila[i_importe], float):
            continue
        fecha = fila[i_fecha] if i_fecha is not None else None
        if isinstance(fecha, float):
            fecha = xlrd.xldate_as_datetime(fecha, libro.datemode).date()
        numero = fila[i_numero] if i_numero is not None else None
        movimientos.append(Movimiento(fecha, str(fila[i_concepto]).strip(), fila[i_importe],
                                      int(numero) if isinstance(numero, float) else numero))
    return movimientos


def leer_conocidos(ruta, familias):
    """pagadores_conocidos.csv: lo que ya has confirmado a mano (pagador del banco -> niño)."""
    conocidos = {}
    if not ruta.exists():
        ruta.write_text("pagador;nino\n", encoding="utf-8-sig")
        return conocidos
    por_nino = {" ".join(tokens(n.original)): fam for fam in familias for n in fam.ninos}
    for fila in csv.DictReader(ruta.read_text(encoding="utf-8-sig").splitlines(), delimiter=";"):
        pagador, nino = fila.get("pagador", ""), fila.get("nino", "")
        familia = por_nino.get(" ".join(tokens(nino)))
        if not pagador.strip():
            continue
        if familia is None:
            print(f"  Aviso: en {ruta.name}, el niño '{nino}' no está en el listado de inscritos.")
            continue
        conocidos[frozenset(tokens(pagador))] = familia
    return conocidos


# ---------------------------------------------------------------- emparejar

def nivel_adulto(adulto, pagador):
    """Compara un padre/madre del listado con el nombre que trae el banco."""
    if not adulto.apellidos:
        return None, ""
    nombre = any(contiene(n, pagador) for n in adulto.nombres)
    apellidos = sum(contiene(a, pagador) for a in adulto.apellidos)
    if nombre and apellidos == len(adulto.apellidos) and len(adulto.apellidos) >= 2:
        return "ALTA", f"coincide con {adulto.original}"
    if nombre and apellidos:
        return "MEDIA", f"coincide con {adulto.original} (faltan apellidos en el listado)"
    if apellidos >= 2:
        return "MEDIA", f"coinciden los dos apellidos de {adulto.original}, pero no el nombre"
    if apellidos:
        return "BAJA", f"comparte apellido con {adulto.original}, ¿familiar?"
    return None, ""


def nivel_familia(familia, pagador):
    mejor = (None, "")
    for adulto in familia.padres:
        nivel, motivo = nivel_adulto(adulto, pagador)
        if nivel and (mejor[0] is None or NIVELES.index(nivel) < NIVELES.index(mejor[0])):
            mejor = (nivel, motivo)
    if mejor[0] is None:
        # Pista débil: el niño suele llevar el primer apellido de cada progenitor.
        for nino in familia.ninos:
            comunes = [a for a in nino.apellidos if contiene(a, pagador)]
            if comunes:
                return "BAJA", f"comparte apellido {comunes[0].title()} con {nino.original}"
    return mejor


@dataclass
class Resultado:
    movimiento: Movimiento
    nivel: str
    familia: Familia = None
    motivo: str = ""
    alternativas: list = field(default_factory=list)


def emparejar(movimiento, familias, conocidos):
    pagador = set(tokens(movimiento.concepto))
    familia = conocidos.get(frozenset(pagador))
    if familia:
        return Resultado(movimiento, "CONFIRMADA", familia, "confirmado en pagadores_conocidos.csv")

    candidatos = []
    for familia in familias:
        nivel, motivo = nivel_familia(familia, pagador)
        if nivel:
            candidatos.append((NIVELES.index(nivel), nivel, familia, motivo))
    if not candidatos:
        return Resultado(movimiento, "SIN IDENTIFICAR")
    candidatos.sort(key=lambda c: c[0])
    _, nivel, familia, motivo = candidatos[0]
    empatados = [c for c in candidatos if c[0] == candidatos[0][0]]
    otras = [f"{c[2].etiqueta} ({c[1].lower()})" for c in candidatos[1:4]]
    if len(empatados) > 1:
        return Resultado(movimiento, "AMBIGUA", None,
                         f"{len(empatados)} familias con confianza {nivel.lower()}",
                         [f"{c[2].etiqueta} — {c[3]}" for c in empatados[:4]])
    return Resultado(movimiento, nivel, familia, motivo, otras)


# ---------------------------------------------------------------- informe

def escribir_informe(ruta, resultados, familias, gastos):
    libro = openpyxl.Workbook()
    negrita = Font(bold=True)

    def hoja(titulo, cabecera, filas, anchos, columna_nivel=None, primera=False):
        h = libro.active if primera else libro.create_sheet()
        h.title = titulo
        h.append(cabecera)
        for celda in h[1]:
            celda.font = negrita
        for fila in filas:
            h.append(fila)
            if columna_nivel is not None:
                nivel = fila[columna_nivel]
                h.cell(h.max_row, columna_nivel + 1).fill = PatternFill("solid", fgColor=COLORES[nivel])
        for i, ancho in enumerate(anchos):
            h.column_dimensions[chr(ord("A") + i)].width = ancho
        for fila in h.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        h.freeze_panes = "A2"
        h.auto_filter.ref = h.dimensions
        return h

    hoja("Pagos",
         ["Fecha", "Pagador (banco)", "Importe", "Confianza", "Familia sugerida", "Niños", "Motivo", "Otras opciones"],
         [[r.movimiento.fecha, r.movimiento.concepto, r.movimiento.importe, r.nivel,
           r.familia.etiqueta if r.familia else "", r.familia.nombres_ninos if r.familia else "",
           r.motivo, "\n".join(r.alternativas)] for r in resultados],
         [11, 34, 9, 15, 40, 34, 45, 50], columna_nivel=3, primera=True)

    por_familia = {}
    for r in resultados:
        if r.familia and r.nivel in ("CONFIRMADA", "ALTA", "MEDIA"):
            por_familia.setdefault(r.familia.clave, []).append(r.movimiento.importe)
    hoja("Por familia",
         ["Familia", "Niños", "Sección", "Nº pagos", "Total recibido", "Importes", "Ya marcado en el listado"],
         [[f.etiqueta, f.nombres_ninos, f.secciones, len(por_familia.get(f.clave, [])),
           sum(por_familia.get(f.clave, [])),
           " + ".join(f"{i:g}" for i in por_familia.get(f.clave, [])), f.marcas] for f in familias],
         [45, 40, 9, 9, 14, 22, 50])

    hoja("Gastos (ignorados)", ["Fecha", "Concepto", "Importe"],
         [[g.fecha, g.concepto, g.importe] for g in gastos], [11, 40, 10])

    for h in libro.worksheets:
        for fila in h.iter_rows(min_row=2):
            if hasattr(fila[0].value, "year"):
                fila[0].number_format = "DD/MM/YYYY"
    libro.save(ruta)


def main():
    parser = argparse.ArgumentParser(description="Sugiere a qué familia pertenece cada ingreso del banco.")
    parser.add_argument("movimientos", help="Excel de movimientos descargado de Unicaja (.xls)")
    parser.add_argument("inscritos", help="Listado de inscritos (.xlsx, .xls o .csv) con padres separados por /")
    parser.add_argument("-o", "--salida", default="conciliacion.xlsx")
    parser.add_argument("--col-padres", help="Columna de padres (letra o número) si no hay cabecera")
    parser.add_argument("--col-nino", help="Columna del niño (letra o número) si no hay cabecera")
    args = parser.parse_args()

    familias = leer_inscritos(args.inscritos, columna(args.col_padres), columna(args.col_nino))
    conocidos = leer_conocidos(Path(__file__).with_name(FICHERO_CONOCIDOS), familias)
    movimientos = leer_movimientos(args.movimientos)
    ingresos = [m for m in movimientos if m.importe > 0]
    gastos = [m for m in movimientos if m.importe <= 0]

    resultados = [emparejar(m, familias, conocidos) for m in ingresos]
    escribir_informe(args.salida, resultados, familias, gastos)

    print(f"{len(familias)} familias ({sum(len(f.ninos) for f in familias)} niños), "
          f"{len(ingresos)} ingresos, {len(gastos)} gastos ignorados.")
    for nivel in NIVELES:
        n = sum(r.nivel == nivel for r in resultados)
        if n:
            print(f"  {nivel:<16} {n:>3}")
    print(f"Informe: {args.salida}")


if __name__ == "__main__":
    main()
